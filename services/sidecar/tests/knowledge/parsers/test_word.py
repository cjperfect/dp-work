"""验证 DOCX 顺序、标题、样式、列表、表格与内部接口。"""
import asyncio
import base64
import io
import zipfile

import pytest
from docx import Document

from src import config
from src.transport.handlers import dispatch, OperationError
from src.knowledge.parsers.word import parse_word
from src.knowledge.parsers import dispatcher, word


def sample_word():
    document = Document()
    document.add_heading("项目说明", level=1)
    paragraph = document.add_paragraph("正文 ")
    paragraph.add_run("重点").bold = True
    document.add_paragraph("条目一", style="List Bullet")
    document.add_paragraph("顺序一", style="List Number")
    table = document.add_table(rows=2, cols=2)
    for cells, values in zip(table.rows, [["名称", "数值"], ["Alpha", "42"]]):
        for cell, value in zip(cells.cells, values):
            cell.text = value
    document.add_paragraph("表格之后的结论")
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_word_structure_order_and_inline_styles():
    result = parse_word(sample_word())
    assert result["parser"] == "python-docx"
    assert result["content"].startswith("# 项目说明")
    assert "**重点**" in result["content"] and "- 条目一" in result["content"]
    assert "1. 顺序一" in result["content"]
    assert result["content"].index("Alpha") < result["content"].index("表格之后的结论")
    assert [block["type"] for block in result["blocks"]] == ["heading", "paragraph", "bulletListItem", "numberedListItem", "table", "paragraph"]
    assert result["blocks"][1]["content"][1]["styles"] == {"bold": True}
    assert result["blocks"][4]["content"]["rows"][1]["cells"] == ["Alpha", "42"]
    assert all("page" not in item for item in result["metadata"])


def test_invalid_empty_and_unpacked_limits(monkeypatch):
    with pytest.raises(ValueError, match="损坏"):
        parse_word(b"not docx")
    stream = io.BytesIO()
    Document().save(stream)
    with pytest.raises(ValueError, match="未提取到正文"):
        parse_word(stream.getvalue())
    with monkeypatch.context() as patch:
        patch.setattr(word, "MAX_UNPACKED_BYTES", 8)
        with pytest.raises(ValueError, match="解压后内容过大"):
            parse_word(sample_word())
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("other.xml", "not Word")
    with pytest.raises(ValueError, match="不是有效"):
        parse_word(stream.getvalue())


def test_docx_process_and_rpc_dispatch(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "TEMP_DIR", tmp_path)
    async def case():
        result = await dispatcher.parse_document(sample_word(), "docx")
        assert result["parser"] == "python-docx" and "Alpha" in result["content"]
        assert not list((tmp_path / "parsing").iterdir())
        response = await dispatch("knowledge.parse", {"file_type": "docx", "content_base64": base64.b64encode(sample_word()).decode()}, None)
        assert response["blocks"][4]["type"] == "table"
        with pytest.raises(OperationError) as error:
            await dispatch("knowledge.parse", {"file_type": "docx", "content_base64": base64.b64encode(b"invalid").decode()}, None)
        assert error.value.status == 422
    asyncio.run(case())
